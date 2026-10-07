#!/usr/bin/env python3
"""
Agente de sanitização de vendas Porsche.

Lê uma planilha .xlsx com as colunas brutas do Schema.md, aplica as regras de
sanitização e grava um novo .xlsx com:
  - todas as colunas originais preservadas;
  - cada coluna sanitizada logo após a sua coluna de origem (inclui, além do
    Schema.md, CustomerNameSanitized e SalespersonSanitized);
  - uma aba "Relatorio" com os checks de qualidade e as divergências em
    relação a colunas sanitizadas pré-existentes (se houver).

Uso:
    python porsche_sanitizer.py ENTRADA.xlsx [SAIDA.xlsx]
"""
import re
import sys
import unicodedata
import calendar
import difflib
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INVALID = "INVALID"

# ----------------------------------------------------------------------------
# Mapa fonte -> saída (Schema.md)
# ----------------------------------------------------------------------------
SANITIZED = {
    "sale_date": "SaleDateSanitized",
    "porsche_model": "PorscheModelSanitized",
    "model_year": "ModelYearSanitized",
    "sale_price": "SalesPriceSanitized",
    "vehicle_mileage": "VehicleMileageSanitized",
    "payment_method": "PayMethodSanitized",
    "city": "CitySanitized",
    "state": "StateSanitized",
    "delivery_status": "DeliveryStatusSanitized",
    # extensão além do Schema.md: colunas de texto livre que estavam sem tratamento
    "customer_name": "CustomerNameSanitized",
    "salesperson": "SalespersonSanitized",
}
EXTRA_COLUMNS = {"CustomerNameSanitized", "SalespersonSanitized"}

# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------
def is_blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def clean(v):
    """Texto com espaços colapsados."""
    return re.sub(r"\s+", " ", str(v)).strip()


def round_half_up(x):
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


# --- números por extenso (inglês) -------------------------------------------
UNITS = {w: i for i, w in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen "
    "fourteen fifteen sixteen seventeen eighteen nineteen".split())}
TENS = {w: 10 * (i + 2) for i, w in enumerate(
    "twenty thirty forty fifty sixty seventy eighty ninety".split())}
SCALES = {"thousand": 1000, "million": 1_000_000}


def words_to_number(text):
    """'eighty two thousand' -> 82000. Retorna None se não for número por extenso."""
    toks = [t for t in re.split(r"[\s\-]+", text.lower()) if t and t != "and"]
    if not toks:
        return None
    total = current = 0
    seen = False
    for t in toks:
        if t in UNITS:
            current += UNITS[t]
        elif t in TENS:
            current += TENS[t]
        elif t == "hundred":
            current = max(current, 1) * 100
        elif t in SCALES:
            total += max(current, 1) * SCALES[t]
            current = 0
        else:
            return None
        seen = True
    return total + current if seen else None


def parse_number(s):
    """
    Interpreta números com separadores mistos:
      '$153,200.50' -> 153200.50 | '$89.750,00' -> 89750.00
      '112.750' (milhar) -> 112750 | '1,5' -> 1.5 | '12,500' -> 12500
    """
    s = re.sub(r"[^\d.,]", "", s)
    if not s or not re.search(r"\d", s):
        return None
    has_dot, has_com = "." in s, "," in s
    if has_dot and has_com:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        grp = "," if dec == "." else "."
        s = s.replace(grp, "").replace(dec, ".")
    elif has_dot or has_com:
        sep = "." if has_dot else ","
        parts = s.split(sep)
        if len(parts) > 2 or len(parts[-1]) == 3:      # separador de milhar
            s = "".join(parts)
        else:                                           # separador decimal
            s = parts[0] + "." + parts[1]
    try:
        return Decimal(s)
    except Exception:
        return None


# ----------------------------------------------------------------------------
# Datas
# ----------------------------------------------------------------------------
MONTHS = {}
for i in range(1, 13):
    MONTHS[calendar.month_name[i].lower()] = i
    MONTHS[calendar.month_abbr[i].lower()] = i
MONTHS["sept"] = 9


def _mk_date(y, m, d):
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return INVALID


def sanitize_date(v):
    if is_blank(v):
        return INVALID
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    s = clean(v)
    # YYYY-MM-DD | YYYY/MM/DD | YYYY.MM.DD
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        return _mk_date(int(m[1]), int(m[2]), int(m[3]))
    # MM/DD/YYYY | MM/DD/YY | MM-DD-YY | MM-DD-YYYY
    m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2}|\d{4})", s)
    if m:
        y = int(m[3])
        if len(m[3]) == 2:
            y += 2000
        return _mk_date(y, int(m[1]), int(m[2]))
    # Month DDth, YYYY | Mon DDth YYYY
    m = re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s)
    if m and m[1].lower() in MONTHS:
        return _mk_date(int(m[3]), MONTHS[m[1].lower()], int(m[2]))
    return INVALID


# ----------------------------------------------------------------------------
# Modelos
# ----------------------------------------------------------------------------
CANONICAL_MODELS = """
911 Carrera|911 Carrera S|911 Carrera GTS|911 Turbo|911 Turbo S|911 GT3|911 GT3 RS|911 Dakar
911 Targa 4|911 Targa 4S|718 Cayman|718 Cayman S|718 Cayman GT4 RS|718 Boxster|718 Boxster GTS
718 Spyder RS|Cayenne|Cayenne S|Cayenne Coupe|Cayenne E-Hybrid|Cayenne Turbo|Cayenne Turbo GT
Macan|Macan S|Macan T|Macan GTS|Macan Electric|Panamera|Panamera 4|Panamera 4S|Panamera Turbo
Panamera Turbo S|Panamera 4 E-Hybrid|Taycan|Taycan 4S|Taycan GTS|Taycan Turbo|Taycan Turbo S
Taycan Cross Turismo
""".replace("\n", "|").split("|")
CANONICAL_MODELS = [m.strip() for m in CANONICAL_MODELS if m.strip()]


def _key(s):
    return re.sub(r"[^a-z0-9]", "", strip_accents(str(s)).lower())


MODEL_BY_KEY = {_key(m): m for m in CANONICAL_MODELS}
UPPER_TOKENS = {"GT", "GTS", "RS", "S", "T", "GT2", "GT3", "GT4", "E"}


def smart_title(text):
    """Title case que preserva siglas/trims (GT3, RS, 4S, E-Hybrid...)."""
    out = []
    for tok in clean(text).split(" "):
        parts = []
        for p in tok.split("-"):
            if p.upper() in UPPER_TOKENS or re.search(r"\d", p):
                parts.append(p.upper())
            else:
                parts.append(p.capitalize())
        out.append("-".join(parts))
    return " ".join(out)


def sanitize_model(v):
    if is_blank(v):
        return INVALID
    key = _key(v)
    if key in MODEL_BY_KEY:
        return MODEL_BY_KEY[key]
    return smart_title(v)            # modelo desconhecido: title case, nunca descartar


# ----------------------------------------------------------------------------
# Ano do modelo
# ----------------------------------------------------------------------------
def _year_ok(y):
    return str(y) if 1990 <= y <= 2035 else INVALID


def sanitize_year(v):
    if is_blank(v):
        return INVALID
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _year_ok(int(v)) if float(v).is_integer() else INVALID
    s = clean(v)
    if re.fullmatch(r"\d{4}(\.0+)?", s):
        return _year_ok(int(float(s)))
    m = re.fullmatch(r"(\d{2})[\s\-](\d{2})", s)                    # 20-24 / 20 24
    if m:
        return _year_ok(int(m[1] + m[2]))
    low = s.lower().replace("-", " ")
    if "thousand" in low:                                          # two thousand twenty two
        n = words_to_number(low)
        return _year_ok(n) if n is not None else INVALID
    toks = low.split()                                             # twenty twenty four
    if len(toks) >= 2 and toks[0] in ("nineteen", "twenty"):
        century = 19 if toks[0] == "nineteen" else 20
        rest = toks[1:]
        if rest[0] in ("oh", "o"):                                 # twenty oh five
            rest = rest[1:]
            r = words_to_number(" ".join(rest)) if rest else None
            valid = r is not None and 0 <= r <= 9
        else:
            r = words_to_number(" ".join(rest))
            valid = r is not None and 0 <= r <= 99
        if valid:
            return _year_ok(century * 100 + r)
    return INVALID


# ----------------------------------------------------------------------------
# Preço
# ----------------------------------------------------------------------------
def sanitize_price(v):
    if is_blank(v):
        return INVALID
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return f"{Decimal(str(v)):.2f}" if v > 0 else INVALID
    s = clean(v)
    low = re.sub(r"\b(usd|us\$|dollars?|dollar)\b", " ", s.lower()).replace("$", " ").strip()
    low = clean(low)
    # número por extenso
    n = words_to_number(low)
    if n is not None:
        return f"{Decimal(n):.2f}" if n > 0 else INVALID
    # sufixo k / m
    m = re.fullmatch(r"([\d.,]+)\s*([km])", low)
    if m:
        num = parse_number(m[1])
        if num is None:
            return INVALID
        mult = 1000 if m[2] == "k" else 1_000_000
        return f"{(num * mult):.2f}" if num > 0 else INVALID
    if re.fullmatch(r"[\d.,\s]+", low):
        num = parse_number(low.replace(" ", ""))
        if num is not None and num > 0:
            return f"{num:.2f}"
    return INVALID


# ----------------------------------------------------------------------------
# Milhagem
# ----------------------------------------------------------------------------
KM_TO_MI = Decimal("0.621371")


def sanitize_mileage(v):
    if is_blank(v):
        return INVALID
    if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
        return str(round_half_up(v)) if v >= 0 else INVALID
    s = clean(v).lower()
    if re.fullmatch(r"(zero( miles?)?|new( car)?|0 ?mi\.?)", s):
        return "0"
    is_km = bool(re.search(r"(?<![a-z])(km|kms|kilomet(er|re)s?)(?![a-z])", s))
    body = re.sub(r"(?<![a-z])(miles?|mi|km|kms|kilomet(er|re)s?)(?![a-z])\.?", " ", s)
    body = clean(body.replace(":", " "))
    n = words_to_number(body) if re.search(r"[a-z]", body) else None
    if n is None:
        n = parse_number(body) if re.fullmatch(r"[\d.,\s]+", body) else None
    if n is None:
        return INVALID
    n = Decimal(str(n))
    if is_km:
        n = n * KM_TO_MI
    return str(round_half_up(n))


# ----------------------------------------------------------------------------
# Forma de pagamento
# ----------------------------------------------------------------------------
def sanitize_payment(v):
    if is_blank(v):
        return INVALID
    s = clean(v).lower()
    k = re.sub(r"[^a-z]", "", s)
    if re.search(r"\bach\b", re.sub(r"[_\-]", " ", s)) or k.startswith("ach"):
        return "ACH Payment"
    if "wire" in k:
        return "Wire Transfer"
    if "banktransfer" in k:
        return "Bank Transfer"
    if "credit" in k:
        return "Credit Card"
    if "debit" in k:
        return "Debit Card"
    if "financ" in k:
        return "Financing"
    if "leas" in k:
        return "Lease"
    if "cash" in k:
        return "Cash"
    if "crypto" in k:
        return "Crypto Payment"
    return smart_title(re.sub(r"[_\-]+", " ", s))


# ----------------------------------------------------------------------------
# Cidade
# ----------------------------------------------------------------------------
ABBREV = {"st": "St.", "st.": "St.", "ft": "Ft.", "ft.": "Ft.", "mt": "Mt.", "mt.": "Mt."}


def sanitize_city(v):
    if is_blank(v):
        return INVALID
    words = []
    for w in clean(v).split(" "):
        lw = w.lower()
        if lw in ABBREV:
            words.append(ABBREV[lw])
        else:
            words.append("-".join(p.capitalize() for p in w.split("-")))
    return " ".join(words)


# ----------------------------------------------------------------------------
# Nomes de pessoas (cliente / vendedor)
# ----------------------------------------------------------------------------
NAME_PARTICLES = {"de", "da", "do", "dos", "das", "del", "della", "di", "van", "von",
                  "der", "den", "la", "le", "bin", "al"}


def _cap_name_word(w):
    parts = []
    for hy in w.split("-"):
        sub = []
        for ap in hy.replace("\u2019", "'").split("'"):
            ap = ap.lower()
            if re.fullmatch(r"mc[a-z]{2,}", ap):                  # McDonald
                ap = "Mc" + ap[2:].capitalize()
            else:
                ap = ap.capitalize()
            sub.append(ap)
        parts.append("'".join(sub))
    return "-".join(parts)


def sanitize_person_name(v):
    """
    'SOPHIA Miller' -> 'Sophia Miller' | 'AMANDA  scott' -> 'Amanda Scott'
    'Daniel-Jones' (sem espaço, 1 hífen) -> 'Daniel Jones'
    Nome com apenas um termo ('kevin') é só capitalizado: não dá para inventar sobrenome.
    """
    if is_blank(v):
        return INVALID
    s = re.sub(r"[\d_]", " ", str(v))                               # números/underscore = ruído
    s = re.sub(r"[^\w\s'\u2019.\-]", " ", s)                         # pontuação de ruído (!!, ?, etc.)
    s = clean(s)
    if not re.search(r"[^\W\d_]", s):
        return INVALID
    if " " not in s and s.count("-") == 1 and all(len(p) >= 2 for p in s.split("-")):
        s = s.replace("-", " ")                                  # separador digitado como hífen
    words = []
    for i, w in enumerate(s.split(" ")):
        if i > 0 and w.lower() in NAME_PARTICLES:
            words.append(w.lower())
        else:
            words.append(_cap_name_word(w))
    return " ".join(words)


# ----------------------------------------------------------------------------
# Estado
# ----------------------------------------------------------------------------
STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "florida": "FL", "georgia": "GA",
    "hawaii": "HI", "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA",
    "kansas": "KS", "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV", "new hampshire": "NH",
    "new jersey": "NJ", "new mexico": "NM", "new york": "NY", "north carolina": "NC",
    "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR", "pennsylvania": "PA",
    "rhode island": "RI", "south carolina": "SC", "south dakota": "SD", "tennessee": "TN",
    "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY", "district of columbia": "DC",
}
STATE_CODES = set(STATES.values())


def sanitize_state(v):
    if is_blank(v):
        return INVALID
    s = clean(v).replace(".", "").lower()
    if s.upper() in STATE_CODES:
        return s.upper()
    return STATES.get(s, INVALID)


# ----------------------------------------------------------------------------
# Status de entrega
# ----------------------------------------------------------------------------
DELIVERY_LABELS = ["Delivered", "Pending", "In Transit", "Cancelled", "Awaiting Delivery",
                   "Awaiting Pickup", "Pending Approval", "Pending Review", "Shipped",
                   "Awaiting Review"]
DELIVERY_BY_KEY = {_key(l): l for l in DELIVERY_LABELS}
DELIVERY_BY_KEY.update({"canceled": "Cancelled", "deliverd": "Delivered"})


def sanitize_delivery(v):
    if is_blank(v):
        return INVALID
    k = _key(v)
    if k in DELIVERY_BY_KEY:
        return DELIVERY_BY_KEY[k]
    close = difflib.get_close_matches(k, DELIVERY_BY_KEY, n=1, cutoff=0.88)
    return DELIVERY_BY_KEY[close[0]] if close else INVALID


FUNCS = {
    "sale_date": sanitize_date,
    "porsche_model": sanitize_model,
    "model_year": sanitize_year,
    "sale_price": sanitize_price,
    "vehicle_mileage": sanitize_mileage,
    "payment_method": sanitize_payment,
    "city": sanitize_city,
    "state": sanitize_state,
    "delivery_status": sanitize_delivery,
    "customer_name": sanitize_person_name,
    "salesperson": sanitize_person_name,
}

# ----------------------------------------------------------------------------
# Pipeline
# ----------------------------------------------------------------------------
HDR_RAW = PatternFill("solid", fgColor="1F2937")     # coluna original
HDR_SAN = PatternFill("solid", fgColor="065F46")     # coluna sanitizada (Schema.md)
HDR_EXT = PatternFill("solid", fgColor="1D4ED8")     # coluna sanitizada extra (nomes)
HDR_REP = PatternFill("solid", fgColor="374151")
BAND = PatternFill("solid", fgColor="F3F4F6")
FONT = "Arial"
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# alinhamento por coluna: centro (códigos/datas/ids), direita (números), esquerda (texto)
CENTER_COLS = {"sale_id", "sale_date", "SaleDateSanitized", "model_year", "ModelYearSanitized",
               "state", "StateSanitized", "DeliveryStatusSanitized"}
RIGHT_COLS = {"sale_price", "SalesPriceSanitized", "vehicle_mileage", "VehicleMileageSanitized"}


def read_source(path):
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    existing_san = [h for h in header if h in SANITIZED.values()]
    raw_cols = [h for h in header if h not in existing_san]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(is_blank(x) for x in r):
            continue
        rows.append(dict(zip(header, r)))
    return ws.title, header, raw_cols, rows


def build_layout(raw_cols):
    layout = []
    for c in raw_cols:
        layout.append(c)
        if c in SANITIZED:
            layout.append(SANITIZED[c])
    return layout


def style_header(cell, fill):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def run(src, dst):
    sheet_name, header, raw_cols, rows = read_source(src)
    missing = [c for c in SANITIZED if c not in raw_cols]
    if missing:
        raise SystemExit(f"Colunas obrigatórias ausentes: {missing}")
    layout = build_layout(raw_cols)

    # ---- sanitização (+ comparação com colunas sanitizadas pré-existentes)
    diffs = []
    for r in rows:
        for src_col, out_col in SANITIZED.items():
            new = FUNCS[src_col](r.get(src_col))
            old = r.get(out_col)
            if old is not None and str(old) != new:
                diffs.append((r.get("sale_id"), out_col, r.get(src_col), old, new))
            r[out_col] = new

    # ---- escrita da aba de dados
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.row_dimensions[1].height = 30
    for j, col in enumerate(layout, 1):
        fill = HDR_EXT if col in EXTRA_COLUMNS else HDR_SAN if col in SANITIZED.values() else HDR_RAW
        style_header(ws.cell(row=1, column=j, value=col), fill)

    for i, r in enumerate(rows, 2):
        for j, col in enumerate(layout, 1):
            v = r.get(col)
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FONT)
            c.border = BORDER
            if i % 2 == 1:
                c.fill = BAND
            h = "center" if col in CENTER_COLS else "right" if col in RIGHT_COLS else "left"
            c.alignment = Alignment(horizontal=h, vertical="center")
            if isinstance(v, (datetime, date)):
                c.number_format = "yyyy-mm-dd"
            elif col == "sale_id" and isinstance(v, (int, float)):
                c.number_format = "0"

    last_row = len(rows) + 1
    last_col = get_column_letter(len(layout))
    # INVALID em vermelho via formatação condicional (acompanha edições manuais)
    ws.conditional_formatting.add(
        f"A2:{last_col}{last_row}",
        CellIsRule(operator="equal", formula=['"INVALID"'],
                   font=Font(name=FONT, bold=True, color="9B1C1C"),
                   fill=PatternFill("solid", bgColor="FDE2E2", fgColor="FDE2E2")))

    for j, col in enumerate(layout, 1):
        longest = max([len(str(col))] + [len(str(r.get(col, ""))) for r in rows])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 12), 34)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    ws.sheet_view.zoomScale = 100

    # ---- checks de qualidade
    n = len(rows)
    iso = re.compile(r"\d{4}-\d{2}-\d{2}")
    ids = [r.get("sale_id") for r in rows]
    checks = [
        ("Colunas originais preservadas", all(c in layout for c in raw_cols)),
        ("Colunas sanitizadas logo após a origem",
         all(layout.index(o) == layout.index(s) + 1 for s, o in SANITIZED.items())),
        ("Datas em ISO ou INVALID",
         all(iso.fullmatch(r["SaleDateSanitized"]) or r["SaleDateSanitized"] == INVALID for r in rows)),
        ("Estados USPS (2 letras) ou INVALID",
         all(r["StateSanitized"] in STATE_CODES or r["StateSanitized"] == INVALID for r in rows)),
        ("Milhagem em milhas inteiras",
         all(re.fullmatch(r"\d+", r["VehicleMileageSanitized"]) or r["VehicleMileageSanitized"] == INVALID for r in rows)),
        ("Preço com duas casas decimais",
         all(re.fullmatch(r"\d+\.\d{2}", r["SalesPriceSanitized"]) or r["SalesPriceSanitized"] == INVALID for r in rows)),
        ("Nenhum valor em branco nas colunas sanitizadas",
         all(not is_blank(r[o]) for r in rows for o in SANITIZED.values())),
        ("sale_id único (sem duplicatas)", len(ids) == len(set(ids))),
    ]

    rp = wb.create_sheet("Relatorio")
    rp.append(["Check de qualidade", "Resultado"])
    for name, ok in checks:
        rp.append([name, "OK" if ok else "FALHOU"])
    rp.append([])
    rp.append(["Registros processados", n])
    rp.append([])
    rp.append(["Coluna sanitizada", "Total INVALID"])
    for o in SANITIZED.values():
        rp.append([o, sum(1 for r in rows if r[o] == INVALID)])
    rp.append([])
    rp.append(["Nota", "CustomerNameSanitized e SalespersonSanitized são extensões além do Schema.md "
                       "(Title Case, espaços colapsados; 'Nome-Sobrenome' sem espaço vira 'Nome Sobrenome')."])
    rp.append([])
    rp.append([f"Divergências vs. colunas sanitizadas pré-existentes: {len(diffs)}"])
    if diffs:
        rp.append(["sale_id", "coluna", "valor bruto", "valor anterior", "valor novo"])
        for d in diffs:
            rp.append([d[0], d[1], str(d[2]), str(d[3]), d[4]])
    for row in rp.iter_rows():
        for c in row:
            c.font = Font(name=FONT)
            if c.value in ("Check de qualidade", "Resultado", "Coluna sanitizada", "Total INVALID",
                           "sale_id", "coluna", "valor bruto", "valor anterior", "valor novo"):
                style_header(c, HDR_REP)
            elif c.value == "OK":
                c.font = Font(name=FONT, bold=True, color="065F46")
                c.alignment = Alignment(horizontal="center")
            elif c.value == "FALHOU":
                c.font = Font(name=FONT, bold=True, color="9B1C1C")
                c.alignment = Alignment(horizontal="center")
    rp.column_dimensions["A"].width = 52
    for col in "BCDE":
        rp.column_dimensions[col].width = 26

    wb.save(dst)
    return n, checks, diffs


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + "_tratado.xlsx")
    n, checks, diffs = run(src, dst)
    print(f"{n} registros processados -> {dst}")
    for name, ok in checks:
        print(("[OK]    " if ok else "[FALHA] ") + name)
    print(f"Divergências vs. sanitização pré-existente: {len(diffs)}")
    for d in diffs:
        print("  ", d)